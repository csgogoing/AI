"""用例台账 Excel (tests/generated/test_cases.xlsx) 生成与更新。

列: 脚本序号 | 脚本名称 | 用例标题 | 前置条件 | 测试步骤 | 预期结果 | 测试结果 | 失败详情
数据源: Script/*.py + Script/*.meta.json
结果回填: result_map = {func: (测试结果, 失败详情)}, 来自执行结果与 pytest 失败摘要。
"""
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = ["脚本序号", "脚本名称", "用例标题", "前置条件", "测试步骤", "预期结果", "测试结果", "失败详情"]
COL_WIDTHS = [10, 34, 28, 28, 42, 32, 10, 46]
CENTER_COLS = {1, 7}


def build_ledger(script_dir: Path, xlsx_path: Path,
                 result_map: dict | None = None) -> Path:
    """扫描 Script 目录生成/更新用例台账 Excel。

    result_map: {func: (测试结果, 失败详情)}; 缺失的用例记为"未执行"。
    """
    result_map = result_map or {}
    rows: list[list[str]] = []
    for script in sorted(script_dir.glob("test_smoke_*.py")):
        m = re.match(r"test_smoke_(\d+)", script.name)
        seq = m.group(1) if m else "-"
        meta_file = script.with_suffix(".meta.json")
        cases = []
        if meta_file.exists():
            cases = json.loads(meta_file.read_text(encoding="utf-8")).get("cases", [])
        if not cases:
            rows.append([seq, script.name, "", "", "", "", "未执行", ""])
            continue
        for c in cases:
            func = c.get("func", "")
            result, detail = result_map.get(func, ("未执行", ""))
            rows.append([
                seq, script.name,
                c.get("title", ""), c.get("precondition", ""),
                c.get("steps", ""), c.get("expected", ""),
                result, detail,
            ])

    wb = Workbook()
    ws = wb.active
    ws.title = "用例台账"

    header_fill = PatternFill("solid", fgColor="1F3B57")
    header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    bottom_line = Side(style="thin", color="B8C7D9")
    for col, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=bottom_line)

    center_align = Alignment(horizontal="center", vertical="center")
    wrap_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = center_align if c in CENTER_COLS else wrap_align

    for i, w in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    wb.save(xlsx_path)
    return xlsx_path
